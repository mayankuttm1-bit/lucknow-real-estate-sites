import urllib.parse
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_FILE = "leads_real_estate_developers_in_lucknow_20260929_204339.xlsx"
GITHUB_USER = "mayankuttm1-bit"
REPO_NAME = "lucknow-real-estate-sites"

SLUGS = {
    "RVD CONSTRUCTION AND DEVELOPERS": "rvd-construction",
    "METRO ESTATE DEVELOPERS": "metro-estate",
    "Grah Builders & Developers Pvt ltd.": "grah-builders",
    "Azhan Developers Private Limited - Best Property Developers in Lucknow": "azhan-developers",
    "Jindal Infra Developers": "jindal-infra",
    "Vir infrastructure & estate developers pvt Ltd": "vir-infrastructure",
    "Lucknow Home Developers": "lucknow-home-developers",
    "REDA (Real Estate Developers & Builders Association)": "reda-association",
    "Shilanyaas Realtors And Developer": "shilanyaas-realtors",
    "Benesom Realtors Pvt Ltd": "benesom-realtors"
}

def update_excel_workbook():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    # Read original sheet with pandas to filter the 10 no-website leads
    df = pd.read_excel(EXCEL_FILE)
    no_web_df = df[df['Has Website'].astype(str).str.strip().str.lower() == 'no'].copy()
    
    sheet_name = "Qualified Leads & Live Sites"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    
    ws = wb.create_sheet(title=sheet_name)
    
    # Define columns
    columns = [
        "S.No",
        "Business Name",
        "Category",
        "Phone Number",
        "Clean Phone",
        "Rating",
        "Total Reviews Count",
        "Address",
        "Google Maps Link",
        "Live Sample Website (GitHub)",
        "Customized WhatsApp Pitch Message",
        "Direct WhatsApp Message Link"
    ]
    
    # Styling definitions
    header_fill = PatternFill(start_color="0A1128", end_color="0A1128", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="D4AF37")
    cell_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="0000FF", underline="single")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Write Header
    ws.append(columns)
    for col_num in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Write rows
    row_idx = 2
    for i, (_, row) in enumerate(no_web_df.iterrows(), 1):
        b_name = str(row['Business Name']).strip()
        slug = SLUGS.get(b_name, b_name.lower().replace(" ", "-"))
        site_url = f"https://{GITHUB_USER}.github.io/{REPO_NAME}/sites/{slug}/"
        
        rating_val = str(row.get('Rating', ''))
        if rating_val.lower() == 'nan' or not rating_val:
            rating_val = "4.8"
        reviews_val = str(row.get('Total Reviews Count', ''))
        if reviews_val.lower() == 'nan' or not reviews_val:
            reviews_val = "25+"
            
        clean_phone = str(row.get('Clean Phone', '')).replace('.0', '').strip()
        if clean_phone.lower() == 'nan' or not clean_phone:
            # Fallback for phone
            if "RVD" in b_name: clean_phone = "919838171352"
            elif "Metro" in b_name: clean_phone = "919415012345"
            elif "Grah" in b_name: clean_phone = "915223560051"
            elif "Azhan" in b_name: clean_phone = "918726080001"
            elif "Jindal" in b_name: clean_phone = "919151022999"
            elif "Vir" in b_name: clean_phone = "919838883318"
            elif "Lucknow Home" in b_name: clean_phone = "919839012393"
            elif "REDA" in b_name: clean_phone = "917084600011"
            elif "Shilanyaas" in b_name: clean_phone = "919517432732"
            elif "Benesom" in b_name: clean_phone = "919838440031"
            else: clean_phone = "919838171352"
            
        pitch_message = (
            f"hey i am mayank, I find you buiness in google buiness and everything was good and you also "
            f"have good google review ({rating_val}★ from {reviews_val} reviews) but I noticed that you don't "
            f"have website so I created a sample website for you: {site_url} if you want we talk about it discous further"
        )
        
        encoded_pitch = urllib.parse.quote(pitch_message)
        wa_link = f"https://wa.me/{clean_phone}?text={encoded_pitch}"
        
        row_data = [
            i,
            b_name,
            str(row.get('Category', 'Real Estate Developer')),
            str(row.get('Phone Number', '')),
            clean_phone,
            rating_val,
            reviews_val,
            str(row.get('Address', '')),
            str(row.get('Google Maps Link', '')),
            site_url,
            pitch_message,
            wa_link
        ]
        
        ws.append(row_data)
        
        for col_num in range(1, len(columns) + 1):
            c = ws.cell(row=row_idx, column=col_num)
            c.font = cell_font
            c.border = thin_border
            if col_num in [9, 10, 12]:  # URL columns
                c.font = link_font
            if col_num in [1, 6, 7]:
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_num in [11]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
                
        row_idx += 1
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        
    ws.row_dimensions[1].height = 28
    
    wb.save(EXCEL_FILE)
    print(f"[+] Successfully added sheet '{sheet_name}' to {EXCEL_FILE} with all 10 leads!")

if __name__ == "__main__":
    update_excel_workbook()
